import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from plotly.subplots import make_subplots
import openpyxl
from datetime import datetime
import io
import os

# ── PDF imports ──────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, PageBreak, HRFlowable)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.graphics import renderPDF
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
from reportlab.platypus import Image as RLImage
import tempfile

# ═══════════════════════════════════════════════════════════════════════════
# PAGE CONFIG
# ═══════════════════════════════════════════════════════════════════════════
st.set_page_config(
    page_title="Energy Intelligence · Sapphire NeelaBlue",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ═══════════════════════════════════════════════════════════════════════════
# THEME / CSS
# ═══════════════════════════════════════════════════════════════════════════
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Syne:wght@400;600;700;800&family=DM+Sans:wght@300;400;500&display=swap');

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif;
}

/* ── Background ── */
.stApp {
    background: #0a0e1a;
    color: #e8eaf0;
}

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background: #0d1220 !important;
    border-right: 1px solid #1e2a42;
}
[data-testid="stSidebar"] .stMarkdown p,
[data-testid="stSidebar"] label {
    color: #8892a4 !important;
    font-size: 0.78rem;
    letter-spacing: 0.08em;
    text-transform: uppercase;
}
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 {
    color: #e8eaf0 !important;
    font-family: 'Syne', sans-serif;
}

/* ── Metric cards ── */
[data-testid="stMetric"] {
    background: #111827;
    border: 1px solid #1e2a42;
    border-radius: 12px;
    padding: 18px 20px !important;
    position: relative;
    overflow: hidden;
}
[data-testid="stMetric"]::before {
    content: '';
    position: absolute;
    top: 0; left: 0; right: 0;
    height: 2px;
    background: linear-gradient(90deg, #00d4ff, #7c3aed);
}
[data-testid="stMetricLabel"] {
    color: #8892a4 !important;
    font-size: 0.72rem !important;
    letter-spacing: 0.1em;
    text-transform: uppercase;
}
[data-testid="stMetricValue"] {
    color: #e8eaf0 !important;
    font-family: 'Syne', sans-serif !important;
    font-size: 1.6rem !important;
    font-weight: 700 !important;
}
[data-testid="stMetricDelta"] {
    font-size: 0.75rem !important;
}

/* ── Section headers ── */
.section-header {
    font-family: 'Syne', sans-serif;
    font-size: 1.05rem;
    font-weight: 700;
    color: #e8eaf0;
    letter-spacing: 0.04em;
    text-transform: uppercase;
    margin: 0 0 4px 0;
    padding-bottom: 8px;
    border-bottom: 1px solid #1e2a42;
}

/* ── Page title ── */
.page-title {
    font-family: 'Syne', sans-serif;
    font-size: 2rem;
    font-weight: 800;
    color: #e8eaf0;
    letter-spacing: -0.02em;
    line-height: 1.1;
}
.page-subtitle {
    font-family: 'DM Sans', sans-serif;
    font-size: 0.85rem;
    color: #4a5568;
    margin-top: 2px;
    letter-spacing: 0.06em;
    text-transform: uppercase;
}

/* ── Accent badge ── */
.badge {
    display: inline-block;
    padding: 2px 10px;
    border-radius: 20px;
    font-size: 0.7rem;
    font-weight: 600;
    letter-spacing: 0.08em;
    text-transform: uppercase;
}
.badge-blue  { background: #0d2240; color: #00d4ff; border: 1px solid #0a3a5c; }
.badge-green { background: #0d2a1a; color: #00e676; border: 1px solid #0a3a20; }
.badge-amber { background: #2a1a05; color: #ffab40; border: 1px solid #3a2a0a; }

/* ── Plotly chart container ── */
.chart-card {
    background: #111827;
    border: 1px solid #1e2a42;
    border-radius: 14px;
    padding: 20px;
    margin-bottom: 16px;
}

/* ── Selectbox / slider ── */
[data-testid="stSelectbox"] > div,
[data-testid="stMultiSelect"] > div {
    background: #111827 !important;
    border-color: #1e2a42 !important;
    border-radius: 8px !important;
    color: #e8eaf0 !important;
}
.stSlider > div { color: #e8eaf0; }

/* ── Number input ── */
[data-testid="stNumberInput"] input {
    background: #111827 !important;
    border-color: #1e2a42 !important;
    color: #e8eaf0 !important;
    border-radius: 8px !important;
}

/* ── Buttons ── */
.stButton > button {
    background: linear-gradient(135deg, #0d2a4a, #1a1040) !important;
    border: 1px solid #1e3a5a !important;
    color: #00d4ff !important;
    font-family: 'Syne', sans-serif !important;
    font-weight: 600 !important;
    letter-spacing: 0.06em !important;
    border-radius: 8px !important;
    padding: 10px 20px !important;
    transition: all 0.2s !important;
}
.stButton > button:hover {
    border-color: #00d4ff !important;
    box-shadow: 0 0 16px rgba(0,212,255,0.15) !important;
}

/* ── Download button ── */
[data-testid="stDownloadButton"] > button {
    background: linear-gradient(135deg, #003d2a, #001a3a) !important;
    border: 1px solid #00e676 !important;
    color: #00e676 !important;
    font-family: 'Syne', sans-serif !important;
    font-weight: 600 !important;
    letter-spacing: 0.06em !important;
    border-radius: 8px !important;
    width: 100% !important;
}
[data-testid="stDownloadButton"] > button:hover {
    box-shadow: 0 0 16px rgba(0,230,118,0.2) !important;
}

/* ── Divider ── */
hr { border-color: #1e2a42 !important; }

/* ── Tabs ── */
[data-testid="stTabs"] [data-testid="stTab"] {
    font-family: 'Syne', sans-serif !important;
    font-weight: 600 !important;
    color: #4a5568 !important;
    letter-spacing: 0.05em;
    font-size: 0.8rem;
    text-transform: uppercase;
}
[data-testid="stTabs"] [aria-selected="true"] {
    color: #00d4ff !important;
    border-bottom-color: #00d4ff !important;
}

/* ── Dataframe ── */
[data-testid="stDataFrame"] {
    border: 1px solid #1e2a42;
    border-radius: 10px;
    overflow: hidden;
}

/* ── Scrollbar ── */
::-webkit-scrollbar { width: 6px; }
::-webkit-scrollbar-track { background: #0a0e1a; }
::-webkit-scrollbar-thumb { background: #1e2a42; border-radius: 3px; }
</style>
""", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════════════════════
# PLOTLY TEMPLATE
# ═══════════════════════════════════════════════════════════════════════════
PLOT_LAYOUT = dict(
    paper_bgcolor='#111827',
    plot_bgcolor='#111827',
    font=dict(family='DM Sans', color='#8892a4', size=11),
    xaxis=dict(gridcolor='#1e2a42', linecolor='#1e2a42', tickfont=dict(color='#8892a4')),
    yaxis=dict(gridcolor='#1e2a42', linecolor='#1e2a42', tickfont=dict(color='#8892a4')),
    legend=dict(bgcolor='#0d1220', bordercolor='#1e2a42', borderwidth=1,
                font=dict(color='#8892a4', size=10)),
    margin=dict(l=50, r=20, t=40, b=50),
    hovermode='x unified',
)

LT_COLORS = {
    'Rebeaming LT':              '#00d4ff',
    'Sizing LT':                 '#7c3aed',
    'Weaving LT-1':              '#00e676',
    'Weaving LT-2':              '#ffab40',
    'Weaving LT-3':              '#ff6b6b',
    'FINISHING LT':              '#f06292',
    'Finishing LT':              '#f06292',
    'Utility  BOILER LT':        '#80cbc4',
    'Utility AIR COMPRESSOR LT': '#ce93d8',
    'Utility AIR COMPRESSOR LT-2': '#fff176',
}

# ═══════════════════════════════════════════════════════════════════════════
# DATA LOADING
# ═══════════════════════════════════════════════════════════════════════════
@st.cache_data
def load_data(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
    records = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        date_str = sheet_name.strip()
        try:
            date = datetime.strptime(date_str, "%d-%m-%Y")
        except:
            continue

        rows = list(ws.iter_rows(values_only=True))
        current_lt = None
        lt_main = None
        lt_solar = None

        for row in rows:
            # LT section header: col0 is non-empty string, col1 is None, not a header row
            if (row[0] is not None and row[1] is None
                    and isinstance(row[0], str)
                    and len(row[0].strip()) > 2
                    and row[0].strip() not in ('Sr .#',)):
                # Save previous LT before switching
                if current_lt is not None and lt_main is not None:
                    records.append({
                        'date': date, 'lt': current_lt,
                        'main_kwh': max(lt_main, 0),
                        'solar_kwh': max(lt_solar or 0, 0)
                    })
                current_lt = row[0].strip()
                lt_main = None
                lt_solar = None
                continue

            if current_lt is None:
                continue

            # Main Breaker: always calculate from col D - col E (idx 3, 4)
            if row[1] == 'Main  Breaker':
                try:
                    t, y = row[3], row[4]
                    lt_main = (t - y) if isinstance(t, (int,float)) and isinstance(y, (int,float)) else 0
                except:
                    lt_main = 0

            # Solar row
            if row[1] is not None and 'Solar' in str(row[1]):
                try:
                    t, y = row[3], row[4]
                    lt_solar = (t - y) if isinstance(t, (int,float)) and isinstance(y, (int,float)) else 0
                except:
                    lt_solar = 0

        # Save last LT of sheet
        if current_lt is not None and lt_main is not None:
            records.append({
                'date': date, 'lt': current_lt,
                'main_kwh': max(lt_main, 0),
                'solar_kwh': max(lt_solar or 0, 0)
            })

    df = pd.DataFrame(records)
    df['lt'] = df['lt'].str.strip()
    df.loc[df['lt'] == 'Finishing LT', 'lt'] = 'FINISHING LT'
    df['main_kwh'] = df['main_kwh'].clip(lower=0)
    df['solar_kwh'] = df['solar_kwh'].clip(lower=0)
    df = df.sort_values('date').reset_index(drop=True)
    return df

# ── File uploader ──
uploaded_file = st.sidebar.file_uploader(
    "📂 Upload Load Report (.xlsx)",
    type=["xlsx"],
    help="Upload your Daily Load Report Excel file"
)

if uploaded_file is None:
    st.markdown("""
    <div style="display:flex;flex-direction:column;align-items:center;justify-content:center;
                min-height:60vh;text-align:center;">
        <div style="background:#111827;border:1px solid #1e2a42;border-radius:18px;
                    padding:56px 64px;max-width:520px;">
            <p style="font-size:2.8rem;margin:0 0 16px;">⚡</p>
            <p style="font-family:Syne,sans-serif;font-size:1.6rem;color:#e8eaf0;
                      font-weight:800;margin:0 0 10px;letter-spacing:-0.02em;">
                Energy Intelligence</p>
            <p style="color:#4a5568;font-size:0.8rem;letter-spacing:0.1em;
                      text-transform:uppercase;margin:0 0 24px;">
                Sapphire Fibres · NeelaBlue Denim Unit</p>
            <div style="background:#0d1220;border:1px dashed #1e3a5a;border-radius:10px;
                        padding:20px 24px;margin-bottom:0;">
                <p style="color:#00d4ff;font-size:0.85rem;margin:0 0 6px;font-weight:600;">
                    📂 Upload your Daily Load Report</p>
                <p style="color:#8892a4;font-size:0.78rem;margin:0;">
                    Use the uploader in the sidebar to get started.<br>
                    Supports any month's Excel report.
                </p>
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)
    st.stop()

df = load_data(uploaded_file.read())
ALL_LTS = sorted(df['lt'].unique().tolist())

# ═══════════════════════════════════════════════════════════════════════════
# SIDEBAR
# ═══════════════════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown('<p class="page-title" style="font-size:1.3rem;">⚡ NeelaBlue</p>', unsafe_allow_html=True)
    st.markdown('<p class="page-subtitle">Energy Intelligence Dashboard</p>', unsafe_allow_html=True)
    st.markdown("---")

    st.markdown("**TARIFF SETTINGS**")
    blended_rate = st.number_input("Blended Rate — All Sources (Rs/kWh)", value=22.38, step=0.01, format="%.2f",
                                    help="Blended cost across LESCO + Solar + Gas Engine")
    lesco_rate   = st.number_input("LESCO-Only Rate (Rs/kWh)", value=46.00, step=0.01, format="%.2f",
                                    help="Pure LESCO tariff — used to calculate savings vs grid-only")
    st.markdown("---")

    st.markdown("**FILTERS**")
    selected_lts = st.multiselect("Select LT Substations", ALL_LTS, default=ALL_LTS,
                                   help="Filter which LTs to display")

    date_range = st.slider("Date Range (Day of May)",
                           min_value=1, max_value=31, value=(1, 31))
    st.markdown("---")

    st.markdown("**PDF REPORT**")
    pdf_lts = st.multiselect("LTs to include in Report", ALL_LTS, default=ALL_LTS)
    pdf_title = st.text_input("Report Title", value="Energy Report – May 2026")

# ═══════════════════════════════════════════════════════════════════════════
# FILTER DATA
# ═══════════════════════════════════════════════════════════════════════════
start_day = datetime(2026, 5, date_range[0])
end_day   = datetime(2026, 5, date_range[1])

fdf = df[
    (df['lt'].isin(selected_lts)) &
    (df['date'] >= start_day) &
    (df['date'] <= end_day)
].copy()

fdf['actual_cost']   = fdf['main_kwh'] * blended_rate
fdf['lesco_cost']    = fdf['main_kwh'] * lesco_rate
fdf['solar_savings'] = fdf['solar_kwh'] * (lesco_rate - blended_rate)
fdf['solar_offset_pct'] = fdf.apply(
    lambda r: round(r['solar_kwh'] / r['main_kwh'] * 100, 1) if r['main_kwh'] > 0 else 0, axis=1)

# ═══════════════════════════════════════════════════════════════════════════
# HEADER
# ═══════════════════════════════════════════════════════════════════════════
col_title, col_badge = st.columns([3, 1])
with col_title:
    st.markdown('<p class="page-title">Energy Intelligence</p>', unsafe_allow_html=True)
    st.markdown('<p class="page-subtitle">Sapphire Fibres · NeelaBlue Denim Unit · May 2026</p>', unsafe_allow_html=True)
with col_badge:
    st.markdown(f"""
    <div style="text-align:right; padding-top:10px;">
        <span class="badge badge-blue">⚡ {len(selected_lts)} LTs Active</span><br><br>
        <span class="badge badge-green">☀ Solar Monitored</span>
    </div>""", unsafe_allow_html=True)

st.markdown("<br>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════════════════════
# KPI CARDS
# ═══════════════════════════════════════════════════════════════════════════
total_kwh    = fdf['main_kwh'].sum()
total_solar  = fdf['solar_kwh'].sum()
total_cost   = fdf['actual_cost'].sum()
total_savings= fdf['solar_savings'].sum()
avg_offset   = (total_solar / total_kwh * 100) if total_kwh > 0 else 0
peak_day_grp = fdf.groupby('date')['main_kwh'].sum()
peak_day     = peak_day_grp.idxmax().strftime("%-d May") if not peak_day_grp.empty else "—"
peak_kwh     = peak_day_grp.max() if not peak_day_grp.empty else 0

k1, k2, k3, k4, k5 = st.columns(5)
k1.metric("Total Grid Consumption", f"{total_kwh:,.0f} kWh",
          delta=f"{total_kwh/1000:.1f} MWh total")
k2.metric("Total Solar Generated", f"{total_solar:,.0f} kWh",
          delta=f"{avg_offset:.1f}% avg offset")
k3.metric("Actual Cost (Blended)", f"Rs {total_cost/1e6:.2f}M",
          delta=f"Blended: LESCO + Solar + Gas @ Rs {blended_rate}")
k4.metric("Solar Savings vs LESCO", f"Rs {total_savings/1e6:.2f}M",
          delta=f"Savings vs LESCO-only @ Rs {lesco_rate}")
k5.metric("Peak Consumption Day", peak_day,
          delta=f"{peak_kwh:,.0f} kWh")

st.markdown("<br>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════════════════════
# TABS
# ═══════════════════════════════════════════════════════════════════════════
tab1, tab2, tab3, tab4 = st.tabs(["📈  TRENDS", "☀  SOLAR ANALYSIS", "💰  COST IMPACT", "📋  DATA TABLE"])

# ─── TAB 1: TRENDS ───────────────────────────────────────────────────────
with tab1:
    st.markdown('<p class="section-header">Daily Load Trend — All LTs</p>', unsafe_allow_html=True)

    daily_lt = fdf.groupby(['date', 'lt'])['main_kwh'].sum().reset_index()

    fig1 = go.Figure()
    for lt in selected_lts:
        sub = daily_lt[daily_lt['lt'] == lt]
        if sub.empty: continue
        c = LT_COLORS.get(lt, '#aaaaaa')
        fig1.add_trace(go.Scatter(
            x=sub['date'], y=sub['main_kwh'],
            name=lt, mode='lines',
            line=dict(color=c, width=2),
            fill='tozeroy',
            fillcolor=c.replace('#', 'rgba(') + ',0.06)' if c.startswith('#') else c,
            hovertemplate=f"<b>{lt}</b><br>%{{x|%d %b}}: %{{y:,.0f}} kWh<extra></extra>"
        ))
    fig1.update_layout(**PLOT_LAYOUT, height=360,
                       title=dict(text="Grid Consumption by LT (kWh/day)", font=dict(color='#e8eaf0', size=13, family='Syne')))
    st.plotly_chart(fig1, use_container_width=True)

    col_a, col_b = st.columns(2)
    with col_a:
        st.markdown('<p class="section-header">LT-wise Monthly Total (kWh)</p>', unsafe_allow_html=True)
        lt_total = fdf.groupby('lt')['main_kwh'].sum().sort_values(ascending=True).reset_index()
        fig2 = go.Figure(go.Bar(
            x=lt_total['main_kwh'], y=lt_total['lt'],
            orientation='h',
            marker=dict(
                color=[LT_COLORS.get(lt, '#aaaaaa') for lt in lt_total['lt']],
                opacity=0.85,
                line=dict(width=0)
            ),
            hovertemplate="<b>%{y}</b><br>%{x:,.0f} kWh<extra></extra>"
        ))
        fig2.update_layout(**PLOT_LAYOUT, height=320,
                           xaxis_title="kWh", yaxis_title="",
                           title=dict(text="Total Consumption Ranking", font=dict(color='#e8eaf0', size=12, family='Syne')))
        st.plotly_chart(fig2, use_container_width=True)

    with col_b:
        st.markdown('<p class="section-header">Daily Plant Total vs Solar</p>', unsafe_allow_html=True)
        daily_total = fdf.groupby('date').agg(main_kwh=('main_kwh','sum'), solar_kwh=('solar_kwh','sum')).reset_index()
        fig3 = go.Figure()
        fig3.add_trace(go.Bar(x=daily_total['date'], y=daily_total['main_kwh'],
                              name='Grid Consumption', marker_color='#1e3a5a',
                              hovertemplate="%{x|%d %b}<br>Grid: %{y:,.0f} kWh<extra></extra>"))
        fig3.add_trace(go.Bar(x=daily_total['date'], y=daily_total['solar_kwh'],
                              name='Solar Generated', marker_color='#00d4ff',
                              hovertemplate="%{x|%d %b}<br>Solar: %{y:,.0f} kWh<extra></extra>"))
        fig3.update_layout(**PLOT_LAYOUT, barmode='overlay', height=320,
                           title=dict(text="Plant-wide: Grid vs Solar", font=dict(color='#e8eaf0', size=12, family='Syne')))
        st.plotly_chart(fig3, use_container_width=True)

# ─── TAB 2: SOLAR ────────────────────────────────────────────────────────
with tab2:
    st.markdown('<p class="section-header">Solar Generation & Offset Analysis</p>', unsafe_allow_html=True)

    col_s1, col_s2 = st.columns(2)
    with col_s1:
        lt_solar_sum = fdf.groupby('lt')[['main_kwh','solar_kwh']].sum().reset_index()
        lt_solar_sum['offset_pct'] = (lt_solar_sum['solar_kwh'] / lt_solar_sum['main_kwh'] * 100).clip(0, 200).round(1)
        fig_s1 = go.Figure()
        fig_s1.add_trace(go.Bar(x=lt_solar_sum['lt'], y=lt_solar_sum['main_kwh'],
                                name='Grid', marker_color='#1e3a5a'))
        fig_s1.add_trace(go.Bar(x=lt_solar_sum['lt'], y=lt_solar_sum['solar_kwh'],
                                name='Solar', marker_color='#00d4ff'))
        fig_s1.update_layout(**PLOT_LAYOUT, barmode='group', height=340,
                             title=dict(text="Solar vs Grid per LT (May Total)", font=dict(color='#e8eaf0', size=12, family='Syne')),
                             xaxis_tickangle=-30)
        st.plotly_chart(fig_s1, use_container_width=True)

    with col_s2:
        fig_s2 = go.Figure(go.Bar(
            x=lt_solar_sum['lt'],
            y=lt_solar_sum['offset_pct'],
            marker=dict(
                color=lt_solar_sum['offset_pct'],
                colorscale=[[0,'#1e2a42'],[0.5,'#0d4a6a'],[1,'#00d4ff']],
                showscale=False,
                line=dict(width=0)
            ),
            hovertemplate="<b>%{x}</b><br>Solar Offset: %{y:.1f}%<extra></extra>"
        ))
        fig_s2.update_layout(**PLOT_LAYOUT, height=340,
                             title=dict(text="Solar Offset % per LT", font=dict(color='#e8eaf0', size=12, family='Syne')),
                             xaxis_tickangle=-30,
                             yaxis_title="Offset %")
        st.plotly_chart(fig_s2, use_container_width=True)

    st.markdown('<p class="section-header">Daily Solar Offset % by LT</p>', unsafe_allow_html=True)
    daily_offset = fdf.groupby(['date','lt']).agg(
        main=('main_kwh','sum'), solar=('solar_kwh','sum')).reset_index()
    daily_offset['pct'] = (daily_offset['solar'] / daily_offset['main'] * 100).clip(0,200).fillna(0)
    fig_s3 = go.Figure()
    for lt in selected_lts:
        sub = daily_offset[daily_offset['lt'] == lt]
        if sub.empty: continue
        c = LT_COLORS.get(lt, '#aaaaaa')
        fig_s3.add_trace(go.Scatter(
            x=sub['date'], y=sub['pct'], name=lt, mode='lines+markers',
            line=dict(color=c, width=1.5),
            marker=dict(size=4, color=c),
            hovertemplate=f"<b>{lt}</b><br>%{{x|%d %b}}: %{{y:.1f}}%<extra></extra>"
        ))
    fig_s3.add_hline(y=100, line_dash='dash', line_color='#00e676',
                     annotation_text="100% offset", annotation_font_color='#00e676')
    fig_s3.update_layout(**PLOT_LAYOUT, height=340,
                         title=dict(text="Daily Solar Offset % (Solar kWh / Grid kWh)", font=dict(color='#e8eaf0', size=12, family='Syne')),
                         yaxis_title="Offset %")
    st.plotly_chart(fig_s3, use_container_width=True)

# ─── TAB 3: COST ─────────────────────────────────────────────────────────
with tab3:
    st.markdown('<p class="section-header">Cost Impact & Solar Savings</p>', unsafe_allow_html=True)

    col_c1, col_c2 = st.columns(2)
    with col_c1:
        lt_cost = fdf.groupby('lt').agg(
            actual=('actual_cost','sum'),
            lesco=('lesco_cost','sum'),
            savings=('solar_savings','sum')
        ).reset_index().sort_values('actual', ascending=False)
        fig_c1 = go.Figure()
        fig_c1.add_trace(go.Bar(x=lt_cost['lt'], y=lt_cost['lesco']/1000,
                                name='Without Solar/Gas (LESCO-only)', marker_color='#3a1a1a'))
        fig_c1.add_trace(go.Bar(x=lt_cost['lt'], y=lt_cost['actual']/1000,
                                name='Actual (All Sources Blended)', marker_color='#00d4ff'))
        fig_c1.update_layout(**PLOT_LAYOUT, barmode='overlay', height=340,
                             title=dict(text="Actual vs Without-Solar Cost per LT (Rs '000)", font=dict(color='#e8eaf0', size=12, family='Syne')),
                             xaxis_tickangle=-30,
                             yaxis_title="Rs '000")
        st.plotly_chart(fig_c1, use_container_width=True)

    with col_c2:
        fig_c2 = go.Figure(go.Bar(
            x=lt_cost['lt'], y=lt_cost['savings']/1000,
            marker=dict(color='#00e676', opacity=0.8),
            hovertemplate="<b>%{x}</b><br>Savings: Rs %{y:,.1f}K<extra></extra>"
        ))
        fig_c2.update_layout(**PLOT_LAYOUT, height=340,
                             title=dict(text="Solar Savings per LT (Rs '000)", font=dict(color='#e8eaf0', size=12, family='Syne')),
                             xaxis_tickangle=-30,
                             yaxis_title="Rs '000")
        st.plotly_chart(fig_c2, use_container_width=True)

    # Daily cost trend
    daily_cost = fdf.groupby('date').agg(
        actual=('actual_cost','sum'),
        savings=('solar_savings','sum')
    ).reset_index()
    fig_c3 = make_subplots(specs=[[{"secondary_y": True}]])
    fig_c3.add_trace(go.Bar(x=daily_cost['date'], y=daily_cost['actual']/1000,
                            name='Daily Cost (Rs K)', marker_color='#1e3a5a'), secondary_y=False)
    fig_c3.add_trace(go.Scatter(x=daily_cost['date'], y=daily_cost['savings']/1000,
                                name='Solar Savings (Rs K)', mode='lines',
                                line=dict(color='#00e676', width=2)), secondary_y=True)
    fig_c3.update_layout(**PLOT_LAYOUT, height=300,
                         title=dict(text="Daily Cost & Solar Savings Trend", font=dict(color='#e8eaf0', size=12, family='Syne')))
    fig_c3.update_yaxes(title_text="Cost (Rs K)", secondary_y=False, gridcolor='#1e2a42', tickfont=dict(color='#8892a4'))
    fig_c3.update_yaxes(title_text="Savings (Rs K)", secondary_y=True, gridcolor='#1e2a42', tickfont=dict(color='#00e676'))
    st.plotly_chart(fig_c3, use_container_width=True)

# ─── TAB 4: DATA TABLE ───────────────────────────────────────────────────
with tab4:
    st.markdown('<p class="section-header">Raw Data — LT-wise Daily Summary</p>', unsafe_allow_html=True)
    display_df = fdf[['date','lt','main_kwh','solar_kwh','solar_offset_pct','actual_cost','solar_savings']].copy()
    display_df['date'] = display_df['date'].dt.strftime('%d-%b-%Y')
    display_df.columns = ['Date','LT','Grid kWh','Solar kWh','Solar Offset %','Actual Cost (Rs)','Solar Savings (Rs)']
    display_df['Actual Cost (Rs)'] = display_df['Actual Cost (Rs)'].round(0)
    display_df['Solar Savings (Rs)'] = display_df['Solar Savings (Rs)'].round(0)
    st.dataframe(display_df, use_container_width=True, height=500)

st.markdown("<br>", unsafe_allow_html=True)
st.markdown("---")

# ═══════════════════════════════════════════════════════════════════════════
# PDF GENERATION
# ═══════════════════════════════════════════════════════════════════════════
def make_mpl_chart(dates, series_dict, title, ylabel, colors_map, height_in=2.8):
    fig, ax = plt.subplots(figsize=(7.2, height_in), facecolor='#111827')
    ax.set_facecolor('#111827')
    for name, vals in series_dict.items():
        c = colors_map.get(name, '#aaaaaa')
        ax.plot(range(len(dates)), vals, color=c, linewidth=1.5, label=name)
    ax.set_xticks(range(0, len(dates), max(1, len(dates)//8)))
    ax.set_xticklabels([dates[i].strftime('%-d') for i in range(0, len(dates), max(1, len(dates)//8))],
                        color='#8892a4', fontsize=7)
    ax.tick_params(axis='y', colors='#8892a4', labelsize=7)
    ax.set_ylabel(ylabel, color='#8892a4', fontsize=7)
    ax.set_title(title, color='#e8eaf0', fontsize=9, fontweight='bold', pad=6)
    ax.spines[:].set_color('#1e2a42')
    ax.grid(axis='y', color='#1e2a42', linewidth=0.5)
    if len(series_dict) > 1:
        ax.legend(fontsize=6, framealpha=0.3, facecolor='#0d1220', edgecolor='#1e2a42',
                  labelcolor='#8892a4', loc='upper right', ncol=2)
    plt.tight_layout(pad=0.4)
    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='#111827')
    plt.close()
    buf.seek(0)
    return buf

def generate_pdf(report_df, lt_list, title_str, blended, lesco):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    W = A4[0] - 36*mm

    # ── Styles ──
    base = getSampleStyleSheet()
    title_style = ParagraphStyle('ptitle', fontName='Helvetica-Bold',
                                  fontSize=18, textColor=colors.HexColor('#e8eaf0'),
                                  spaceAfter=2, alignment=TA_LEFT)
    sub_style   = ParagraphStyle('psub', fontName='Helvetica',
                                  fontSize=8, textColor=colors.HexColor('#4a5568'),
                                  spaceAfter=12, alignment=TA_LEFT, leading=12)
    h2_style    = ParagraphStyle('ph2', fontName='Helvetica-Bold',
                                  fontSize=11, textColor=colors.HexColor('#00d4ff'),
                                  spaceBefore=14, spaceAfter=4)
    body_style  = ParagraphStyle('pbody', fontName='Helvetica',
                                  fontSize=8.5, textColor=colors.HexColor('#c0c8d8'),
                                  leading=13)
    label_style = ParagraphStyle('plabel', fontName='Helvetica',
                                  fontSize=7.5, textColor=colors.HexColor('#8892a4'),
                                  alignment=TA_CENTER)
    val_style   = ParagraphStyle('pval', fontName='Helvetica-Bold',
                                  fontSize=13, textColor=colors.HexColor('#e8eaf0'),
                                  alignment=TA_CENTER)

    story = []

    # ── Cover Header ──
    story.append(Paragraph(title_str, title_style))
    story.append(Paragraph(
        f"Sapphire Fibres Limited · NeelaBlue Denim Unit · Generated {datetime.now().strftime('%d %b %Y, %H:%M')}",
        sub_style))
    story.append(HRFlowable(width=W, thickness=1, color=colors.HexColor('#1e2a42'), spaceAfter=10))

    # ── Tariff Info ──
    story.append(Paragraph("Tariff Settings", h2_style))
    tariff_data = [
        ['Parameter', 'Value'],
        ['Blended Rate (LESCO + Solar + Gas Engine)', f'Rs {blended:.2f} / kWh'],
        ['LESCO-Only Rate (benchmark)', f'Rs {lesco:.2f} / kWh'],
        ['LTs Covered', ', '.join(lt_list)],
    ]
    t_tariff = Table(tariff_data, colWidths=[W*0.45, W*0.55])
    t_tariff.setStyle(TableStyle([
        ('BACKGROUND',   (0,0), (-1,0),  colors.HexColor('#0d2240')),
        ('TEXTCOLOR',    (0,0), (-1,0),  colors.HexColor('#00d4ff')),
        ('FONTNAME',     (0,0), (-1,0),  'Helvetica-Bold'),
        ('FONTSIZE',     (0,0), (-1,-1), 8),
        ('BACKGROUND',   (0,1), (-1,-1), colors.HexColor('#111827')),
        ('TEXTCOLOR',    (0,1), (-1,-1), colors.HexColor('#c0c8d8')),
        ('GRID',         (0,0), (-1,-1), 0.5, colors.HexColor('#1e2a42')),
        ('ROWBACKGROUNDS',(0,1),(-1,-1), [colors.HexColor('#111827'), colors.HexColor('#0d1220')]),
        ('LEFTPADDING',  (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ('TOPPADDING',   (0,0), (-1,-1), 5),
        ('BOTTOMPADDING',(0,0), (-1,-1), 5),
    ]))
    story.append(t_tariff)
    story.append(Spacer(1, 10))

    # ── Summary KPIs per LT ──
    story.append(Paragraph("LT-wise Summary", h2_style))
    lt_sum = report_df.groupby('lt').agg(
        grid=('main_kwh','sum'),
        solar=('solar_kwh','sum'),
        actual=('actual_cost','sum'),
        savings=('solar_savings','sum')
    ).reset_index()
    lt_sum['offset_pct'] = (lt_sum['solar'] / lt_sum['grid'] * 100).clip(0,200).round(1)

    kpi_data = [['LT Substation', 'Grid (kWh)', 'Solar (kWh)', 'Offset %', 'Cost (Rs)', 'Savings (Rs)']]
    for _, row in lt_sum.iterrows():
        kpi_data.append([
            row['lt'],
            f"{row['grid']:,.0f}",
            f"{row['solar']:,.0f}",
            f"{row['offset_pct']:.1f}%",
            f"{row['actual']:,.0f}",
            f"{row['savings']:,.0f}",
        ])
    total_row = [
        'TOTAL',
        f"{lt_sum['grid'].sum():,.0f}",
        f"{lt_sum['solar'].sum():,.0f}",
        f"{lt_sum['solar'].sum()/lt_sum['grid'].sum()*100:.1f}%",
        f"{lt_sum['actual'].sum():,.0f}",
        f"{lt_sum['savings'].sum():,.0f}",
    ]
    kpi_data.append(total_row)

    col_w = [W*0.28, W*0.13, W*0.13, W*0.1, W*0.18, W*0.18]
    t_kpi = Table(kpi_data, colWidths=col_w)
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND',   (0,0), (-1,0),  colors.HexColor('#0d2240')),
        ('TEXTCOLOR',    (0,0), (-1,0),  colors.HexColor('#00d4ff')),
        ('FONTNAME',     (0,0), (-1,0),  'Helvetica-Bold'),
        ('FONTNAME',     (0,-1),(-1,-1), 'Helvetica-Bold'),
        ('BACKGROUND',   (0,-1),(-1,-1), colors.HexColor('#0d2240')),
        ('TEXTCOLOR',    (0,-1),(-1,-1), colors.HexColor('#00e676')),
        ('FONTSIZE',     (0,0), (-1,-1), 7.5),
        ('BACKGROUND',   (0,1), (-1,-2), colors.HexColor('#111827')),
        ('TEXTCOLOR',    (0,1), (-1,-2), colors.HexColor('#c0c8d8')),
        ('ROWBACKGROUNDS',(0,1),(-1,-2), [colors.HexColor('#111827'), colors.HexColor('#0d1220')]),
        ('GRID',         (0,0), (-1,-1), 0.4, colors.HexColor('#1e2a42')),
        ('ALIGN',        (1,0), (-1,-1), 'RIGHT'),
        ('LEFTPADDING',  (0,0), (-1,-1), 7),
        ('RIGHTPADDING', (0,0), (-1,-1), 7),
        ('TOPPADDING',   (0,0), (-1,-1), 5),
        ('BOTTOMPADDING',(0,0), (-1,-1), 5),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 14))

    # ── Charts per LT ──
    dates_all = sorted(report_df['date'].unique())

    for lt in lt_list:
        story.append(PageBreak())
        story.append(Paragraph(f"LT Detail: {lt}", h2_style))
        story.append(HRFlowable(width=W, thickness=0.5, color=colors.HexColor('#1e2a42'), spaceAfter=8))

        lt_df = report_df[report_df['lt'] == lt].sort_values('date')
        if lt_df.empty:
            story.append(Paragraph("No data for this LT in selected range.", body_style))
            continue

        # Mini KPIs
        g  = lt_df['main_kwh'].sum()
        s  = lt_df['solar_kwh'].sum()
        ac = lt_df['actual_cost'].sum()
        sv = lt_df['solar_savings'].sum()
        op = s/g*100 if g > 0 else 0
        kpi_mini = [
            [Paragraph('Grid Consumed', label_style), Paragraph('Solar Generated', label_style),
             Paragraph('Solar Offset %', label_style), Paragraph('Actual Cost', label_style),
             Paragraph('Solar Savings', label_style)],
            [Paragraph(f"{g:,.0f} kWh", val_style), Paragraph(f"{s:,.0f} kWh", val_style),
             Paragraph(f"{op:.1f}%", val_style), Paragraph(f"Rs {ac:,.0f}", val_style),
             Paragraph(f"Rs {sv:,.0f}", val_style)],
        ]
        t_mini = Table(kpi_mini, colWidths=[W/5]*5)
        t_mini.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#111827')),
            ('BOX',        (0,0), (-1,-1), 0.5, colors.HexColor('#1e2a42')),
            ('INNERGRID',  (0,0), (-1,-1), 0.5, colors.HexColor('#1e2a42')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING',(0,0),(-1,-1),6),
        ]))
        story.append(t_mini)
        story.append(Spacer(1, 10))

        # Grid + Solar chart
        lt_dates = lt_df['date'].tolist()
        grid_vals = lt_df['main_kwh'].tolist()
        solar_vals = lt_df['solar_kwh'].tolist()

        fig_buf = make_mpl_chart(
            lt_dates,
            {'Grid kWh': grid_vals, 'Solar kWh': solar_vals},
            f"{lt} — Daily Grid & Solar (kWh)",
            "kWh",
            {'Grid kWh': '#00d4ff', 'Solar kWh': '#00e676'},
            height_in=2.6
        )
        story.append(RLImage(fig_buf, width=W, height=W*0.36))
        story.append(Spacer(1, 8))

        # Cost chart
        cost_vals = lt_df['actual_cost'].tolist()
        sav_vals  = lt_df['solar_savings'].tolist()
        fig_buf2 = make_mpl_chart(
            lt_dates,
            {'Actual Cost (Rs)': cost_vals, 'Solar Savings (Rs)': sav_vals},
            f"{lt} — Daily Cost & Savings (Rs)",
            "Rs",
            {'Actual Cost (Rs)': '#ffab40', 'Solar Savings (Rs)': '#00e676'},
            height_in=2.6
        )
        story.append(RLImage(fig_buf2, width=W, height=W*0.36))
        story.append(Spacer(1, 10))

        # Daily data table (compact)
        story.append(Paragraph("Daily Breakdown", ParagraphStyle('pdaily', fontName='Helvetica-Bold',
                                fontSize=9, textColor=colors.HexColor('#8892a4'), spaceAfter=4)))
        tbl_data = [['Date', 'Grid kWh', 'Solar kWh', 'Offset %', 'Cost (Rs)', 'Savings (Rs)']]
        for _, r in lt_df.iterrows():
            op_r = r['solar_kwh']/r['main_kwh']*100 if r['main_kwh'] > 0 else 0
            tbl_data.append([
                r['date'].strftime('%d %b'),
                f"{r['main_kwh']:,.0f}",
                f"{r['solar_kwh']:,.0f}",
                f"{op_r:.1f}%",
                f"{r['actual_cost']:,.0f}",
                f"{r['solar_savings']:,.0f}",
            ])
        t_daily = Table(tbl_data, colWidths=[W*0.13, W*0.15, W*0.15, W*0.12, W*0.22, W*0.23])
        t_daily.setStyle(TableStyle([
            ('BACKGROUND',   (0,0), (-1,0),  colors.HexColor('#0d2240')),
            ('TEXTCOLOR',    (0,0), (-1,0),  colors.HexColor('#00d4ff')),
            ('FONTNAME',     (0,0), (-1,0),  'Helvetica-Bold'),
            ('FONTSIZE',     (0,0), (-1,-1), 7),
            ('BACKGROUND',   (0,1), (-1,-1), colors.HexColor('#111827')),
            ('TEXTCOLOR',    (0,1), (-1,-1), colors.HexColor('#c0c8d8')),
            ('ROWBACKGROUNDS',(0,1),(-1,-1), [colors.HexColor('#111827'), colors.HexColor('#0d1220')]),
            ('GRID',         (0,0), (-1,-1), 0.3, colors.HexColor('#1e2a42')),
            ('ALIGN',        (1,0), (-1,-1), 'RIGHT'),
            ('LEFTPADDING',  (0,0), (-1,-1), 5),
            ('RIGHTPADDING', (0,0), (-1,-1), 5),
            ('TOPPADDING',   (0,0), (-1,-1), 3),
            ('BOTTOMPADDING',(0,0), (-1,-1), 3),
        ]))
        story.append(t_daily)

    doc.build(story)
    buf.seek(0)
    return buf

# ── PDF Button ──────────────────────────────────────────────────────────
st.markdown('<p class="section-header">Generate PDF Report</p>', unsafe_allow_html=True)

pdf_col1, pdf_col2 = st.columns([3, 1])
with pdf_col1:
    st.markdown(f"""
    <div style="background:#111827; border:1px solid #1e2a42; border-radius:10px; padding:14px 18px;">
        <span style="color:#8892a4; font-size:0.78rem; letter-spacing:0.08em; text-transform:uppercase;">
        Selected for report:</span><br>
        <span style="color:#e8eaf0; font-family:'Syne',sans-serif; font-weight:600;">
        {len(pdf_lts)} LT{'s' if len(pdf_lts)!=1 else ''} — {pdf_title}</span><br>
        <span style="color:#4a5568; font-size:0.75rem;">
        Blended Rs {blended_rate}/kWh · LESCO Rs {lesco_rate}/kWh · 
        Days {date_range[0]}–{date_range[1]} May 2026</span>
    </div>""", unsafe_allow_html=True)

with pdf_col2:
    if st.button("⚡ Build PDF Report"):
        if not pdf_lts:
            st.error("Select at least one LT for the report.")
        else:
            with st.spinner("Generating report..."):
                pdf_df = df[
                    (df['lt'].isin(pdf_lts)) &
                    (df['date'] >= start_day) &
                    (df['date'] <= end_day)
                ].copy()
                pdf_df['actual_cost']   = pdf_df['main_kwh'] * blended_rate
                pdf_df['lesco_cost']    = pdf_df['main_kwh'] * lesco_rate
                pdf_df['solar_savings'] = pdf_df['solar_kwh'] * (lesco_rate - blended_rate)
                pdf_buf = generate_pdf(pdf_df, pdf_lts, pdf_title, blended_rate, lesco_rate)
            fname = f"energy_report_may2026_{datetime.now().strftime('%H%M%S')}.pdf"
            st.download_button(
                label="📥 Download PDF",
                data=pdf_buf,
                file_name=fname,
                mime="application/pdf"
            )

st.markdown("<br>", unsafe_allow_html=True)
st.markdown(
    '<p style="text-align:center; color:#1e2a42; font-size:0.7rem; letter-spacing:0.1em;">'
    'SAPPHIRE FIBRES · NEELABLUE DENIM UNIT · ENERGY INTELLIGENCE SYSTEM · v1.0</p>',
    unsafe_allow_html=True
)
